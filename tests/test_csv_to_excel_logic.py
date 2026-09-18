from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from tasks.csv_to_excel import logic


def test_dataframe_to_excel_bytes_roundtrips_data():
    df = pd.DataFrame({"name": ["Cafe A", "Cafe B"], "city": ["Amsterdam", "Rotterdam"]})

    excel_bytes = logic.dataframe_to_excel_bytes(df)
    result = pd.read_excel(BytesIO(excel_bytes))

    pd.testing.assert_frame_equal(result, df)


def test_dataframe_to_excel_bytes_uses_given_sheet_name():
    df = pd.DataFrame({"a": [1]})

    excel_bytes = logic.dataframe_to_excel_bytes(df, sheet_name="MySheet")

    wb = load_workbook(BytesIO(excel_bytes))
    assert wb.sheetnames == ["MySheet"]


def test_dataframe_to_excel_bytes_produces_nonempty_xlsx():
    df = pd.DataFrame({"a": [1, 2, 3]})
    excel_bytes = logic.dataframe_to_excel_bytes(df)

    assert len(excel_bytes) > 0
    assert excel_bytes[:2] == b"PK"  # .xlsx is a zip container
